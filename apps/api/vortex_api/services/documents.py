"""Turn uploaded files into text: PDF, DOCX, XLSX, CSV, plain text/code, and images (OCR).

Everything runs locally. OCR uses RapidOCR (ONNX, CPU) when the `ocr` extra is installed.
"""

from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass, field
from pathlib import Path

log = logging.getLogger(__name__)

MAX_CHARS = 2_000_000
MAX_OCR_PAGES = 40
MAX_SHEET_ROWS = 5000

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".json",
    ".csv",
    ".tsv",
    ".log",
    ".xml",
    ".yaml",
    ".yml",
    ".ini",
    ".toml",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".cs",
    ".go",
    ".rs",
    ".php",
    ".rb",
    ".sql",
    ".sh",
    ".ps1",
    ".css",
    ".html",
    ".htm",
    ".kt",
    ".swift",
    ".lua",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}


class UnsupportedFile(ValueError):
    pass


@dataclass
class Extracted:
    pages: list[tuple[int | None, str]] = field(default_factory=list)  # (page number, text)
    ocr: bool = False

    @property
    def text(self) -> str:
        return "\n\n".join(t for _, t in self.pages if t)


def detect_kind(name: str, head: bytes) -> tuple[str, str]:
    """(kind, mime) from the extension, checked against the file's first bytes."""
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        if not head.startswith(b"%PDF"):
            raise UnsupportedFile("O arquivo não parece ser um PDF válido.")
        return "pdf", "application/pdf"
    if ext in (".docx", ".xlsx"):
        if not head.startswith(b"PK\x03\x04"):
            raise UnsupportedFile(f"O arquivo não parece ser um {ext[1:].upper()} válido.")
        return (
            ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            if ext == ".docx"
            else ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        )
    if ext in IMAGE_EXTENSIONS:
        signatures = {
            ".png": [b"\x89PNG"],
            ".jpg": [b"\xff\xd8\xff"],
            ".jpeg": [b"\xff\xd8\xff"],
            ".webp": [b"RIFF"],
            ".bmp": [b"BM"],
        }
        if not any(head.startswith(s) for s in signatures[ext]):
            raise UnsupportedFile("O arquivo não parece ser uma imagem válida.")
        return "image", f"image/{'jpeg' if ext in ('.jpg', '.jpeg') else ext[1:]}"
    if ext in TEXT_EXTENSIONS:
        if b"\x00" in head:
            raise UnsupportedFile("O arquivo parece ser binário, não texto.")
        return ("csv" if ext in (".csv", ".tsv") else "text"), "text/plain"
    raise UnsupportedFile(
        "Tipo de arquivo não suportado. Envie PDF, DOCX, XLSX, CSV, texto/código ou imagem."
    )


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# ---------------------------------------------------------------- OCR


_ocr_engine = None


def ocr_available() -> bool:
    try:
        import rapidocr_onnxruntime  # noqa: F401
    except ImportError:
        return False
    return True


def _ocr_image(image) -> str:
    """OCR a PIL image or numpy array; returns lines in reading order."""
    global _ocr_engine
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR

    if _ocr_engine is None:
        _ocr_engine = RapidOCR()
    result, _ = _ocr_engine(
        np.asarray(image.convert("RGB")) if hasattr(image, "convert") else image
    )
    if not result:
        return ""
    boxes = []
    for box, text, _score in result:
        ys = [p[1] for p in box]
        boxes.append(((min(ys) + max(ys)) / 2, max(ys) - min(ys), min(p[0] for p in box), text))
    boxes.sort(key=lambda b: (b[0], b[2]))
    lines: list[list[tuple[float, str]]] = []
    line_y = None
    for yc, h, x, text in boxes:
        if line_y is None or abs(yc - line_y) > max(h, 1) * 0.6:
            lines.append([])
            line_y = yc
        lines[-1].append((x, text))
    return "\n".join(" ".join(t for _, t in sorted(line)) for line in lines)


# ---------------------------------------------------------------- per format


def _pdf(path: Path) -> Extracted:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    out = Extracted()
    needs_ocr: list[int] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:  # broken page encodings shouldn't kill the whole file
            text = ""
        out.pages.append((i, text))
        if len(text) < 20:
            needs_ocr.append(i)

    # Scanned pages: render and OCR them (if OCR is installed).
    if needs_ocr and ocr_available():
        try:
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(path))
            for i in needs_ocr[:MAX_OCR_PAGES]:
                image = pdf[i - 1].render(scale=2).to_pil()
                text = _ocr_image(image)
                if text:
                    out.pages[i - 1] = (i, text)
                    out.ocr = True
        except Exception as e:
            log.warning("PDF OCR failed for %s: %s", path.name, e)
    return out


def _docx(path: Path) -> Extracted:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return Extracted(pages=[(None, "\n".join(parts))])


def _xlsx(path: Path) -> Extracted:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = Extracted()
    for n, sheet in enumerate(workbook.worksheets, start=1):
        rows = []
        for r, row in enumerate(sheet.iter_rows(values_only=True)):
            if r >= MAX_SHEET_ROWS:
                rows.append(f"... (planilha cortada em {MAX_SHEET_ROWS} linhas)")
                break
            values = ["" if v is None else str(v) for v in row]
            if any(values):
                rows.append(" | ".join(values))
        out.pages.append((n, f"Planilha: {sheet.title}\n" + "\n".join(rows)))
    workbook.close()
    return out


def _csv(data: bytes, delimiter_hint: str) -> Extracted:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = delimiter_hint
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter)):
        if i >= MAX_SHEET_ROWS:
            rows.append(f"... (cortado em {MAX_SHEET_ROWS} linhas)")
            break
        rows.append(" | ".join(row))
    return Extracted(pages=[(None, "\n".join(rows))])


def _image(path: Path) -> Extracted:
    if not ocr_available():
        return Extracted(pages=[(None, "")])
    from PIL import Image

    with Image.open(path) as image:
        image.thumbnail((2400, 2400))
        return Extracted(pages=[(None, _ocr_image(image))], ocr=True)


def extract(path: Path, kind: str) -> Extracted:
    if kind == "pdf":
        out = _pdf(path)
    elif kind == "docx":
        out = _docx(path)
    elif kind == "xlsx":
        out = _xlsx(path)
    elif kind == "csv":
        out = _csv(path.read_bytes(), "\t" if path.suffix.lower() == ".tsv" else ",")
    elif kind == "image":
        out = _image(path)
    else:
        out = Extracted(pages=[(None, _decode(path.read_bytes()))])

    # Normalize and cap the total size.
    total = 0
    capped = []
    for page, text in out.pages:
        text = text.replace("\x00", "").strip()
        if total + len(text) > MAX_CHARS:
            text = text[: MAX_CHARS - total]
        total += len(text)
        capped.append((page, text))
        if total >= MAX_CHARS:
            break
    out.pages = capped
    return out
