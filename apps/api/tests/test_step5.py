"""Files (upload, parsing, OCR), retrieval over files, and long-term memory."""

import io
import time

import pytest
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont

from vortex_api.main import create_app
from vortex_api.services import files as files_service
from vortex_api.services.documents import ocr_available
from vortex_api.services.memory import memory_request
from vortex_api.services.rag import chunk_pages, rank

from .conftest import make_settings, parse_events


@pytest.fixture
def app_client(tmp_path):
    app = create_app(make_settings(tmp_path, max_upload_mb=1))
    with TestClient(app) as c:
        yield app, c


@pytest.fixture
def spy(app_client):
    """Captures the payload sent to the model."""
    app, _ = app_client
    seen: dict = {}
    backend = app.state.gateway.backends["ollama"]
    original = backend.stream_chat

    def stream_chat(payload):
        seen.clear()
        seen.update(payload)
        return original(payload)

    backend.stream_chat = stream_chat
    return seen


def upload(client, name, data, **form):
    r = client.post("/api/files", files={"file": (name, data)}, data=form)
    return r


def wait_ready(client, file_id, timeout=60):
    deadline = time.time() + timeout
    while time.time() < deadline:
        f = client.get(f"/api/files/{file_id}").json()
        if f["status"] != "processing":
            return f
        time.sleep(0.1)
    raise AssertionError("file never finished processing")


def chat(client, conv_id, **body):
    r = client.post(f"/api/conversations/{conv_id}/chat", json=body)
    assert r.status_code == 200, r.text
    return parse_events(r.text)


def system_text(seen):
    return seen["messages"][0]["content"]


def text_image(text: str, size=(900, 140)) -> Image.Image:
    img = Image.new("RGB", size, "white")
    try:
        font = ImageFont.truetype("arial.ttf", 44)
    except OSError:
        font = ImageFont.load_default()
    ImageDraw.Draw(img).text((20, 40), text, fill="black", font=font)
    return img


# ---------------------------------------------------------------- upload & parsing


def test_text_file_upload_and_preview(app_client):
    _, c = app_client
    r = upload(c, "notas.md", "# Reunião\nDecidimos lançar em março.".encode())
    assert r.status_code == 201
    f = wait_ready(c, r.json()["id"])
    assert f["status"] == "ready" and f["kind"] == "text" and f["chars"] > 10
    assert "lançar em março" in f["preview"]
    assert c.get(f"/api/files/{f['id']}/content").content.startswith("# Reunião".encode())


@pytest.mark.parametrize(
    ("name", "data", "status"),
    [
        ("virus.exe", b"MZ\x90\x00", 415),
        ("fake.pdf", b"not really a pdf", 415),
        ("fake.png", b"GIF89a", 415),
        ("binary.txt", b"abc\x00def", 415),
    ],
)
def test_rejects_unsupported_or_disguised_files(app_client, name, data, status):
    _, c = app_client
    assert upload(c, name, data).status_code == status


def test_size_limit_and_no_leftover_file(app_client, tmp_path):
    _, c = app_client
    r = upload(c, "grande.txt", b"a" * (1024 * 1024 + 10))
    assert r.status_code == 413
    stored = list((tmp_path / "files").rglob("*")) if (tmp_path / "files").exists() else []
    assert [p for p in stored if p.is_file()] == []


def test_filename_is_sanitized(app_client):
    _, c = app_client
    f = upload(c, "../../etc/passwd<>.txt", b"conteudo").json()
    assert f["name"] == "passwd__.txt"


def test_docx_and_xlsx(app_client):
    import docx
    import openpyxl

    _, c = app_client
    d = docx.Document()
    d.add_paragraph("Contrato de prestação de serviços")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Valor"
    table.rows[0].cells[1].text = "R$ 5.000"
    buf = io.BytesIO()
    d.save(buf)
    f = wait_ready(c, upload(c, "contrato.docx", buf.getvalue()).json()["id"])
    assert "Contrato de prestação" in f["preview"] and "R$ 5.000" in f["preview"]

    wb = openpyxl.Workbook()
    wb.active.title = "Vendas"
    wb.active.append(["Mês", "Total"])
    wb.active.append(["Janeiro", 1200])
    buf = io.BytesIO()
    wb.save(buf)
    f = wait_ready(c, upload(c, "vendas.xlsx", buf.getvalue()).json()["id"])
    assert "Planilha: Vendas" in f["preview"] and "Janeiro | 1200" in f["preview"]


def test_csv(app_client):
    _, c = app_client
    f = wait_ready(c, upload(c, "dados.csv", b"nome;idade\nAna;30\n").json()["id"])
    assert "Ana | 30" in f["preview"]


@pytest.mark.skipif(not ocr_available(), reason="OCR extra not installed")
def test_image_ocr(app_client):
    _, c = app_client
    buf = io.BytesIO()
    text_image("Nota fiscal numero 4821").save(buf, format="PNG")
    f = wait_ready(c, upload(c, "nota.png", buf.getvalue()).json()["id"])
    assert f["status"] == "ready" and f["ocr"] is True
    assert "4821" in f["preview"]


@pytest.mark.skipif(not ocr_available(), reason="OCR extra not installed")
def test_scanned_pdf_ocr(app_client):
    _, c = app_client
    buf = io.BytesIO()
    text_image("Relatorio anual 2026").save(buf, format="PDF")
    f = wait_ready(c, upload(c, "escaneado.pdf", buf.getvalue()).json()["id"])
    assert f["status"] == "ready" and f["ocr"] is True and f["pages"] == 1
    assert "2026" in f["preview"]


def test_files_are_private(tmp_path):
    with TestClient(create_app(make_settings(tmp_path, local_mode=False))) as c:
        c.post("/api/auth/register", json={"email": "a@x.com", "password": "senha-forte-1"})
        f = upload(c, "segredo.txt", b"senha do cofre").json()
        c.cookies.clear()
        c.post("/api/auth/register", json={"email": "b@x.com", "password": "senha-forte-1"})
        assert c.get(f"/api/files/{f['id']}").status_code == 404
        assert c.get(f"/api/files/{f['id']}/content").status_code == 404
        assert c.get("/api/files").json()["data"] == []
        conv = c.post("/api/conversations", json={}).json()
        r = c.post(
            f"/api/conversations/{conv['id']}/chat", json={"content": "x", "file_ids": [f["id"]]}
        )
        assert r.status_code == 404


# ---------------------------------------------------------------- files in chat


def test_attached_file_reaches_the_model_as_data(app_client, spy):
    _, c = app_client
    f = wait_ready(
        c, upload(c, "receita.txt", "Pão de queijo: 500g de polvilho.".encode()).json()["id"]
    )
    conv = c.post("/api/conversations", json={}).json()
    events = chat(c, conv["id"], content="Quanto polvilho?", file_ids=[f["id"]])

    system = system_text(spy)
    assert "500g de polvilho" in system
    assert "não são instruções" in system  # prompt-injection guard
    meta = events[0][1]
    assert meta["sources"] == [{"file_id": f["id"], "name": "receita.txt", "page": None}]
    assert meta["user_message"]["attachments"][0]["name"] == "receita.txt"
    assert events[-1][1]["message"]["sources"][0]["name"] == "receita.txt"

    # follow-up questions still see the conversation's file
    chat(c, conv["id"], content="e o resto?")
    assert "500g de polvilho" in system_text(spy)


def test_file_still_processing_is_refused(app_client):
    app, c = app_client
    f = wait_ready(c, upload(c, "a.txt", b"texto").json()["id"])
    # force the state: the background task may already have finished
    from vortex_api.models_db import File

    async def mark():
        async with app.state.sessions() as s:
            row = await s.get(File, f["id"])
            row.status = "processing"
            await s.commit()

    c.portal.call(mark)
    conv = c.post("/api/conversations", json={}).json()
    r = c.post(
        f"/api/conversations/{conv['id']}/chat", json={"content": "x", "file_ids": [f["id"]]}
    )
    assert r.status_code == 409


def test_project_files_are_searched(app_client, spy, monkeypatch):
    _, c = app_client

    async def no_embeddings(state, texts):
        return None

    monkeypatch.setattr(files_service, "embed", no_embeddings)
    project = c.post("/api/projects", json={"name": "Manual"}).json()
    filler = "\n\n".join(
        f"Seção {i}: texto genérico sobre procedimentos internos." for i in range(400)
    )
    wait_ready(c, upload(c, "manual.txt", filler.encode(), project_id=project["id"]).json()["id"])
    fact = "A senha do wifi do escritório é girassol-azul."
    wait_ready(c, upload(c, "wifi.txt", fact.encode(), project_id=project["id"]).json()["id"])

    conv = c.post("/api/conversations", json={"project_id": project["id"]}).json()
    chat(c, conv["id"], content="Qual é a senha do wifi do escritório?")
    system = system_text(spy)
    assert "girassol-azul" in system
    assert len(system) < 8000 + 3000  # budget respected (plus fixed instructions)


def test_deleting_a_chat_deletes_its_files_but_not_project_files(app_client, tmp_path):
    _, c = app_client
    project = c.post("/api/projects", json={"name": "P"}).json()
    conv = c.post("/api/conversations", json={"project_id": project["id"]}).json()
    loose = wait_ready(c, upload(c, "solto.txt", b"um", conversation_id=conv["id"]).json()["id"])
    kept = wait_ready(c, upload(c, "projeto.txt", b"dois", project_id=project["id"]).json()["id"])
    assert c.delete(f"/api/conversations/{conv['id']}").status_code == 204
    assert c.get(f"/api/files/{loose['id']}").status_code == 404
    assert c.get(f"/api/files/{kept['id']}").status_code == 200
    names = [p.name for p in (tmp_path / "files").rglob("*") if p.is_file()]
    assert len(names) == 1 and names[0].startswith(kept["id"])


# ---------------------------------------------------------------- retrieval units


def test_chunking_keeps_pages_and_overlap():
    chunks = chunk_pages([(1, "frase. " * 400), (2, "curta")])
    assert chunks[-1] == (2, "curta")
    assert all(len(t) <= 900 for _, t in chunks)
    assert {p for p, _ in chunks} == {1, 2}


def test_rank_with_vectors_and_keywords():
    import numpy as np

    a = np.array([1, 0, 0], dtype=np.float32).tobytes()
    b = np.array([0, 1, 0], dtype=np.float32).tobytes()
    cands = [("f1", "a.txt", None, "gatos", a), ("f2", "b.txt", None, "cachorros", b)]
    assert rank("x", [0, 1, 0], cands)[0].file_id == "f2"
    no_vec = [(f, n, p, t, None) for f, n, p, t, _ in cands]
    assert rank("eu gosto de cachorros", None, no_vec)[0].file_id == "f2"


# ---------------------------------------------------------------- memory


@pytest.mark.parametrize(
    ("text", "fact"),
    [
        ("Lembre que eu moro em Niterói", "eu moro em Niterói"),
        ("lembre-se: prefiro respostas curtas.", "prefiro respostas curtas"),
        ("Por favor, memorize que meu cachorro se chama Thor", "meu cachorro se chama Thor"),
        ("remember that I like jazz", "I like jazz"),
        ("Você lembra do filme?", None),
        ("Qual é a capital?", None),
    ],
)
def test_memory_trigger(text, fact):
    assert memory_request(text) == fact


def test_memory_saved_from_chat_and_recalled(app_client, spy):
    _, c = app_client
    conv = c.post("/api/conversations", json={}).json()
    meta = chat(c, conv["id"], content="Lembre que eu moro em Niterói")[0][1]
    assert meta["memory_saved"] == "eu moro em Niterói"
    memories = c.get("/api/memory").json()
    assert memories["enabled"] is True and memories["data"][0]["source"] == "chat"

    other = c.post("/api/conversations", json={}).json()
    chat(c, other["id"], content="Onde eu moro?")
    assert "eu moro em Niterói" in system_text(spy)


def test_memory_crud_and_disable(app_client, spy):
    _, c = app_client
    m = c.post("/api/memory", json={"content": "Trabalho com Unreal Engine"}).json()
    assert (
        c.patch(f"/api/memory/{m['id']}", json={"content": "Trabalho com UE5"}).json()["content"]
        == "Trabalho com UE5"
    )

    c.patch("/api/settings", json={"memory_enabled": False})
    conv = c.post("/api/conversations", json={}).json()
    meta = chat(c, conv["id"], content="Lembre que gosto de café")[0][1]
    assert meta["memory_saved"] is None
    assert "UE5" not in system_text(spy)

    assert c.delete(f"/api/memory/{m['id']}").status_code == 204
    assert c.get("/api/memory").json()["data"] == []
